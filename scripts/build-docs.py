"""Rebuild editable Word, draw.io and SVG deliverables from versioned sources."""
from pathlib import Path
import re, html, json, xml.etree.ElementTree as ET
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs'

def word():
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2);sec.bottom_margin=Cm(2);sec.left_margin=Cm(2.2);sec.right_margin=Cm(2.2)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        s=doc.styles[name];s.font.name='Microsoft YaHei';s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei');s.font.color.rgb=RGBColor(0,0,0)
        s.font.size=Pt(10.5 if name=='Normal' else 15 if name.startswith('Heading') else 23)
        s.paragraph_format.space_after=Pt(7);s.paragraph_format.line_spacing=1.25
    doc.styles['Title'].paragraph_format.space_after=Pt(14)
    footer=sec.footer.paragraphs[0];footer.alignment=2
    run=footer.add_run('FitStreak  |  ');run.font.size=Pt(8)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    lines=(OUT/'requirements.md').read_text(encoding='utf8').splitlines()
    i=0
    while i<len(lines):
        line=lines[i]
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].startswith('|'):
                parts=[x.strip() for x in lines[i].strip('|').split('|')]
                if not all(re.fullmatch(r'[-:]+',x) for x in parts):rows.append(parts)
                i+=1
            table=doc.add_table(rows=1,cols=len(rows[0]));table.autofit=False
            widths=[2.1,2.0,12.5] if len(rows[0])==3 else [16.6/len(rows[0])]*len(rows[0])
            for n,row in enumerate(rows):
                cells=table.rows[0].cells if n==0 else table.add_row().cells
                for j,text in enumerate(row):
                    cells[j].width=Cm(widths[j]);cells[j].text=text
                    tcPr=cells[j]._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
                    for edge in ['top','left','bottom','right']:
                        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');b.append(el)
                    tcPr.append(b)
                    margin=OxmlElement('w:tcMar')
                    for side in ['top','left','bottom','right']:
                        m=OxmlElement('w:'+side);m.set(qn('w:w'),'90');m.set(qn('w:type'),'dxa');margin.append(m)
                    tcPr.append(margin)
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DDE8E4' if n==0 else 'FFFFFF');tcPr.append(shade)
                    for p in cells[j].paragraphs:
                        p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.12
                        for r in p.runs:r.font.size=Pt(9);r.bold=n==0
                if n==0:
                    repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
                cant=OxmlElement('w:cantSplit');table.rows[n]._tr.get_or_add_trPr().append(cant)
            doc.add_paragraph();continue
        if line.startswith('# '):doc.add_paragraph(line[2:],'Title')
        elif line.startswith('## '):doc.add_heading(line[3:],1)
        elif line.strip():doc.add_paragraph(line)
        i+=1
    doc.core_properties.title='FitStreak 健身打卡小程序需求规格说明书'
    doc.core_properties.subject='需求基线与验收标准';doc.core_properties.author='FitStreak 项目组'
    doc.save(OUT/'FitStreak-需求分析.docx')

def graph(name,title,nodes,edges,width=1120,height=700):
    mx=ET.Element('mxfile',host='app.diagrams.net',version='26.0.0');dia=ET.SubElement(mx,'diagram',id=name,name=title)
    model=ET.SubElement(dia,'mxGraphModel',dx=str(width),dy=str(height),grid='1',page='1',pageWidth=str(width),pageHeight=str(height));root=ET.SubElement(model,'root')
    ET.SubElement(root,'mxCell',id='0');ET.SubElement(root,'mxCell',id='1',parent='0')
    svg=[f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}"><defs><marker id="arrow" markerWidth="10" markerHeight="10" refX="8" refY="3" orient="auto"><path d="M0,0 L0,6 L9,3 z" fill="#72877c"/></marker></defs><rect width="100%" height="100%" fill="#f6f7f3"/><text x="40" y="42" font-family="Microsoft YaHei,Arial" font-size="24" fill="#173d31">{html.escape(title)}</text>']
    lookup={n[0]:n for n in nodes}
    for idx,(src,target,label) in enumerate(edges):
        a,b=lookup[src],lookup[target];sx=a[2]+a[4]/2;sy=a[3]+a[5];tx=b[2]+b[4]/2;ty=b[3]
        if b[2]>a[2]+a[4]:sx=a[2]+a[4];sy=a[3]+a[5]/2;tx=b[2];ty=b[3]+b[5]/2
        cell=ET.SubElement(root,'mxCell',id='e'+str(idx),value=label,edge='1',parent='1',source=src,target=target,style='edgeStyle=orthogonalEdgeStyle;rounded=1;endArrow=block;strokeColor=#72877c;fontSize=12;labelBackgroundColor=#f6f7f3;')
        ET.SubElement(cell,'mxGeometry',relative='1',attrib={'as':'geometry'})
        svg.append(f'<path d="M{sx} {sy} L{tx} {ty}" stroke="#72877c" stroke-width="1.7" fill="none" marker-end="url(#arrow)"/>')
        if label:svg.append(f'<text x="{(sx+tx)/2+5}" y="{(sy+ty)/2-5}" font-family="Microsoft YaHei,Arial" font-size="11" fill="#5a7265">{html.escape(label)}</text>')
    for id,label,x,y,w,h in nodes:
        cell=ET.SubElement(root,'mxCell',id=id,value=label,vertex='1',parent='1',style='rounded=1;whiteSpace=wrap;html=0;fillColor=#ffffff;strokeColor=#cbd9ce;fontColor=#173d31;fontFamily=Microsoft YaHei;fontSize=14;spacing=12;')
        ET.SubElement(cell,'mxGeometry',x=str(x),y=str(y),width=str(w),height=str(h),attrib={'as':'geometry'})
        svg.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="12" fill="#fff" stroke="#cbd9ce"/>')
        for n,line in enumerate(label.split('\n')):svg.append(f'<text x="{x+w/2}" y="{y+27+n*23}" text-anchor="middle" font-family="Microsoft YaHei,Arial" font-size="{15 if n==0 else 12}" fill="#173d31">{html.escape(line)}</text>')
    svg.append('</svg>')
    folder=OUT/'diagrams';folder.mkdir(exist_ok=True)
    ET.ElementTree(mx).write(folder/(name+'.drawio'),encoding='utf-8',xml_declaration=True)
    (folder/(name+'.svg')).write_text('\n'.join(svg),encoding='utf8')
    return mx

def diagrams():
    graph('architecture','FitStreak 系统架构',[
      ('clients','微信小程序 / 响应式 H5\n手机 · 平板 · 桌面',50,90,270,86),('ui','uni-app + Vue 3\n页面组件 · 双语字典 · 请求封装',420,90,290,86),('wx','微信服务\ncode2Session',820,90,240,86),
      ('api','Fastify REST API\nJWT · JSON Schema · 请求限流',420,250,290,90),('domain','领域服务\n计划 · 打卡 · 统计 · 知识 · 用户',420,410,290,90),('db','PostgreSQL 17\n事务 · 外键 · 唯一约束 · 索引',420,575,290,90),
      ('mock','MSW Mock 服务\n正常 / 空态 / 延迟 / 错误',50,250,270,90),('tests','测试与交付\nAPI集成 · Playwright · GitHub CI',50,410,270,90),('devdb','本地开发与测试\nPGlite 隔离数据库',820,575,240,90)
    ],[('clients','ui','共用界面'),('ui','api','HTTPS JSON'),('api','wx','仅服务端密钥'),('api','domain','已验证身份'),('domain','db','参数化 SQL'),('mock','ui','开发环境'),('tests','domain','集成测试'),('domain','devdb','开发配置')])
    graph('technology-selection','FitStreak 技术选型',[
      ('goal','项目约束\n小程序 + H5 · 双语 · 独立后端',360,85,400,82),
      ('front','前端候选\n微信原生 / Taro / uni-app',50,245,280,90),('back','后端候选\nSpring Boot / NestJS / Fastify',420,245,280,90),('data','数据与测试候选\nMySQL / PostgreSQL · MSW',790,245,280,90),
      ('f','采用 uni-app + Vue 3 + TS\n复用跨端组件和业务逻辑',50,420,280,90),('b','采用 Fastify + TypeScript\n契约校验与模块化单体',420,420,280,90),('d','采用 PostgreSQL + MSW\n关系约束 / 可复现异常场景',790,420,280,90),
      ('deliver','工程交付\nOpenAPI · drawio · SQL · Word · Playwright · GitHub',210,590,700,75)
    ],[('goal','front',''),('goal','back',''),('goal','data',''),('front','f','跨端适配'),('back','b','单人维护成本'),('data','d','一致性与可测试性'),('f','deliver',''),('b','deliver',''),('d','deliver','')])
    graph('er','FitStreak 逻辑 E-R 图',[
      ('user','app_user 用户\nPK id\nUK open_id\nnickname · locale · timezone\nweekly_goal · created_at',55,105,300,170),
      ('plan','fitness_plan 计划\nPK id / FK user_id\ntitle · activity · target_minutes\nweekdays · start_date · end_date\narchived_at · created_at',430,105,330,170),
      ('check','checkin 打卡\nPK id / FK (plan_id, user_id)\nUK (user_id, plan_id, local_date)\nduration_minutes · note\nplan_title · activity · created_at',430,385,330,175),
      ('article','article 文章\nPK id\ncategory · reading_minutes\npublished_at',870,105,290,150),
      ('translation','article_translation 译文\nPK (article_id, locale)\nFK article_id\ntitle · summary · body',870,385,290,150)
    ],[('user','plan','1 → 0..N 拥有'),('user','check','1 → 0..N 完成'),('plan','check','1 → 0..N 产生'),('article','translation','1 → 0..N 翻译')],width=1220,height=660)
    (OUT/'diagrams'/'er.mmd').write_text('''erDiagram
 app_user ||--o{ fitness_plan : owns
 app_user ||--o{ checkin : records
 fitness_plan ||--o{ checkin : generates
 article ||--o{ article_translation : translates
 app_user { uuid id PK\n varchar open_id UK\n varchar nickname\n varchar locale\n varchar timezone\n smallint weekly_goal\n timestamptz created_at }
 fitness_plan { uuid id PK\n uuid user_id FK\n varchar title\n varchar activity\n smallint target_minutes\n smallint_array weekdays\n date start_date\n date end_date\n timestamptz archived_at\n timestamptz created_at }
 checkin { uuid id PK\n uuid user_id FK\n uuid plan_id FK\n date local_date\n smallint duration_minutes\n varchar note\n varchar plan_title\n varchar activity\n timestamptz created_at }
 article { uuid id PK\n varchar category\n smallint reading_minutes\n timestamptz published_at }
 article_translation { uuid article_id PK,FK\n varchar locale PK\n varchar title\n varchar summary\n text body }
''',encoding='utf8')
    (OUT/'diagrams'/'architecture.mmd').write_text('''flowchart TB
 C[微信小程序 / 响应式 H5] --> U[uni-app + Vue 3 + TypeScript]
 U -->|HTTPS JSON| A[Fastify REST API · JWT · 参数校验]
 A --> W[微信 code2Session]
 A --> S[计划 / 打卡 / 统计 / 知识 / 用户]
 S --> P[(PostgreSQL 17)]
 S -.开发测试.-> L[(PGlite)]
 M[MSW HTTP Mock] -.替代 API.-> U
 T[Playwright + API集成测试] --> A
''',encoding='utf8')
    (OUT/'diagrams'/'technology-selection.mmd').write_text('''flowchart TB
 R[小程序 + H5 / 双语 / 关系数据 / 单人维护] --> F[原生 / Taro / uni-app]
 R --> B[Spring Boot / NestJS / Fastify]
 R --> D[MySQL / PostgreSQL]
 F --> UF[uni-app + Vue 3 + TypeScript]
 B --> UB[Fastify 模块化单体]
 D --> UD[PostgreSQL + PGlite开发测试]
 UF --> Q[MSW + Playwright + OpenAPI + GitHub CI]
 UB --> Q
 UD --> Q
''',encoding='utf8')

def prototypes():
    folder=ROOT/'design';folder.mkdir(exist_ok=True)
    mx=ET.Element('mxfile',host='app.diagrams.net',version='26.0.0')
    screens=[('overview','首页概览',['问候语与创建计划按钮','运动时长 / 打卡次数 / 连续天数','今日训练：计划名称 + 时长 + 去打卡','最近七天趋势图']),('plans','训练计划',['创建计划按钮','计划卡片：运动项目与时长','训练日与生效日期','编辑 / 归档 / 去打卡']),('plan-form','计划表单',['计划名称（必填）','运动类型 / 时长 1..600','训练日 多选 1..7','开始日期 / 可选结束日期','取消 / 保存']),('checkin','打卡表单',['当前计划名称','实际运动分钟 1..600','感受备注（最多500字）','完成打卡 / 重复错误提示']),('insights','数据统计',['七天 / 三十天切换','累计时长 / 次数 / 运动天数','每日柱状图（零值也展示）','历史列表：名称 / 日期 / 分钟']),('knowledge','知识列表与详情',['分类：全部 / 习惯 / 训练 / 恢复','文章封面 + 标题 + 摘要','点击后展示正文','返回列表 / 中英文同步']),('profile','个人中心',['头像字母 / 昵称','昵称编辑','每周目标 1..7 天','简体中文 / English','保存设置 / 退出登录'])]
    for sid,title,items in screens:
        diagram=ET.SubElement(mx,'diagram',id=sid,name=title);model=ET.SubElement(diagram,'mxGraphModel',page='1',pageWidth='1280',pageHeight='900');root=ET.SubElement(model,'root');ET.SubElement(root,'mxCell',id='0');ET.SubElement(root,'mxCell',id='1',parent='0')
        def cell(id,label,x,y,w,h,fill='#ffffff',font=15,link=None):
            parent=root
            if link:parent=ET.SubElement(root,'UserObject',id=id,label=label,link=link)
            c=ET.SubElement(parent,'mxCell',**({} if link else {'id':id,'value':label}),vertex='1',parent='1',style=f'rounded=1;whiteSpace=wrap;html=0;fillColor={fill};strokeColor=#d4dfcf;fontColor=#264a37;fontSize={font};fontFamily=Microsoft YaHei;')
            ET.SubElement(c,'mxGeometry',x=str(x),y=str(y),width=str(w),height=str(h),attrib={'as':'geometry'})
        cell('title','FitStreak · '+title,35,20,1150,45,'#e8efdf',23)
        cell('phone','',40,95,375,720,'#f6f7f3')
        cell('top','FitStreak                     EN   A',55,112,345,48)
        for i,label in enumerate(items):cell('m'+str(i),label,60,185+i*93,335,74,'#ffffff',14)
        for i,(target,label,_) in enumerate([screens[0],screens[1],screens[4],screens[5],screens[6]]):cell('nav'+str(i),label,55+i*69,750,67,45,'#e8efdf',10,'data:page/id,'+target)
        cell('desktop','桌面布局 1440px → 左侧导航 + 内容网格',475,100,740,55,'#e8efdf',18)
        cell('sidenav','FitStreak\n\n今日概览\n\n训练计划\n\n运动数据\n\n健身知识\n\n个人中心',475,173,160,420,'#ffffff',15)
        for i,label in enumerate(items):cell('d'+str(i),label,655+(i%2)*280,175+(i//2)*140,260,112)
        cell('rules','交互规则：按钮有加载状态；失败保留输入；401返回登录；空态展示创建入口。\n响应式：≤760px 底部导航；>760px 侧栏；375 / 768 / 1440px 验证。\n原型中的控件均可在 diagrams.net 编辑；底部导航关联对应页面。',475,630,740,155,'#f2f5ee',14)
    ET.ElementTree(mx).write(folder/'FitStreak-prototype.drawio',encoding='utf-8',xml_declaration=True)
    (folder/'README.md').write_text('''# FitStreak 原型

`FitStreak-prototype.drawio` 是 diagrams.net 七页可编辑原型：概览、计划、计划表单、打卡表单、统计、知识、个人中心。手机和桌面线框均为可编辑形状，不是截图。底栏包含页内导航链接。

在 https://app.diagrams.net/ 选择设备 → 打开文件。可修改每个控件、尺寸、文本与连线。

`prototype.html` 是可离线打开的点击原型，用于页面结构与基本流程评审；其中数字为明确的示例数据，不连接真实后端。高保真交互与异常状态请运行根目录 README 的 MSW 预览命令。

设计规范：主色 #234C39，背景 #F6F7F3，辅助绿 #E9EFDF，圆角 9/18/22px。字体使用系统无衬线；移动端最大优先使用单列卡片，桌面使用侧栏及三列指标。业务日固定北京时间，显示语言可切换。
''',encoding='utf8')

if __name__=='__main__':
    word();diagrams();prototypes();print('Word document, diagrams and editable prototype generated.')
